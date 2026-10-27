import hashlib
import json
from openpyxl import load_workbook
from hisobdon.db import Store
from hisobdon.ledger import Ledger
from hisobdon.reports import export


def test_ai_export_covers_complete_financial_history_and_distinguishes_period(tmp_path):
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    cash = ledger.create_account("Cash", "UZS", 1000000, "2026-01-01")
    invest = ledger.create_account("Company", "UZS", 500000, "2026-01-01", "investment")
    mid = s.record_message(123, 1, "old original note", "2026-01-03T10:00:00+05:00")
    old = ledger.post(
        dict(kind="expense", account=cash, amount=10000, date="2026-01-03", message_id=mid), "old"
    )
    ledger.cancel(old)
    ledger.post(dict(kind="income", account=cash, amount=20000, date="2026-02-02"), "income")
    ledger.value(invest, 530000, "2026-01-04")
    ledger.value(invest, 550000, "2026-02-01")
    ledger.rate("12800", "2026-01-03")
    s.save_draft(
        "pending", {"mode": "operation", "data": {"note": "unfinished financial note"}, "step": "amount"}
    )
    s.conn.execute("INSERT INTO aliases VALUES (?,?)", ("lunch", "Ovqat"))
    s.set("BOT_TOKEN", "private-test-value")
    s.set("backup_file", "/private/not-for-AI")
    files = export(s, "2026-02-01", "2026-03-01", tmp_path / "exports")
    manifests = [p for p in files if p.name == "manifest.json"]
    assert manifests, "Missing full-archive manifest"
    manifest = json.loads(manifests[0].read_text())
    json_files = [p for p in files if p.suffix == ".json" and p.name != "manifest.json"]
    assert json_files
    bundles = [json.loads(p.read_text()) for p in json_files]
    tables = {}
    for bundle in bundles:
        for name, records in bundle["tables"].items():
            tables.setdefault(name, []).extend(records)
    expected = {
        "accounts",
        "operations",
        "movements",
        "valuations",
        "rates",
        "categories",
        "aliases",
        "messages",
        "drafts",
        "audit",
    }
    assert set(tables) == expected
    for name in expected:
        assert len(tables[name]) == s.one(f"SELECT COUNT(*) AS n FROM {name}")["n"]
    assert len(tables["valuations"]) == 2
    assert tables["operations"][0]["status"] == "cancelled"
    assert tables["messages"][0]["text"] == "old original note"
    assert "chat_id" not in tables["messages"][0]
    assert bundles[0]["metadata"]["requested_period"]["start"] == "2026-02-01"
    assert bundles[0]["analysis"]["period_totals"]["UZS"]["income"] == 20000
    assert bundles[0]["analysis"]["period_totals"]["UZS"]["expense"] == 0
    contents = "".join(p.read_text() for p in files if p.suffix in {".txt", ".json"})
    assert "old original note" in contents and "unfinished financial note" in contents
    assert "private-test-value" not in contents and "/private/not-for-AI" not in contents
    for entry in manifest["files"]:
        path = manifests[0].parent / entry["name"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == entry["sha256"]
    wb = load_workbook(next(p for p in files if p.suffix == ".xlsx"))
    assert (
        "Baholashlar" in wb.sheetnames
        and "Pul harakatlari" in wb.sheetnames
        and "Qoralamalar" in wb.sheetnames
    )


def test_numbered_parts_preserve_every_row_and_long_text(tmp_path, monkeypatch):
    from hisobdon import ai_export

    monkeypatch.setattr(ai_export, "ROWS_PER_PART", 5)
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    ledger.create_account("Reserve", "USD", 50000, "2026-01-01", is_reserve=True)
    long_text = "=this is data, not formula: " + "x" * 40000
    mid = s.record_message(100, 1, long_text, "2026-01-02T10:00:00+05:00")
    files = export(s, "2026-02-01", "2026-03-01", tmp_path / "out")
    bundles = [json.loads(p.read_text()) for p in files if p.suffix == ".json" and p.name != "manifest.json"]
    assert len(bundles) > 1
    messages = [r for b in bundles for r in b["tables"]["messages"]]
    assert len(messages) == 1 and messages[0]["id"] == mid and messages[0]["text"] == long_text
    for name in ai_export.TABLES:
        rows = [r for b in bundles for r in b["tables"][name]]
        assert len(rows) == s.one(f"SELECT count(*) AS n FROM {name}")["n"]
    assert all(b["metadata"]["part_count"] == len(bundles) for b in bundles)
    cells = []
    for path in files:
        if path.suffix == ".xlsx":
            wb = load_workbook(path)
            cells.extend(
                c
                for row in wb["Asl xabarlar"]
                for c in row
                if isinstance(c.value, str) and c.value.startswith("=this")
            )
    assert len(cells) == 1 and cells[0].data_type == "s"
    assert "JSON/TXT" in cells[0].value and len(cells[0].value) <= 32767
    texts = "".join(p.read_text() for p in files if p.suffix == ".txt")
    assert long_text in texts
