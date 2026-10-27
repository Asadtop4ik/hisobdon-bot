import json
from openpyxl import load_workbook
from hisobdon.db import Store
from hisobdon.ledger import Ledger
from hisobdon.reports import totals, export, balance_text
from hisobdon.backup import backup, restore, new_key


def test_summary_does_not_count_principal_or_opening(tmp_path):
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    a = ledger.create_account("A", "UZS", 10000000, "2026-01-01")
    b = ledger.create_account("B", "UZS", 0, "2026-01-01")
    ledger.post(dict(kind="income", account=a, amount=500000, date="2026-01-02"), "i")
    ledger.post(dict(kind="expense", account=a, amount=100000, category="Ovqat", date="2026-01-03"), "e")
    ledger.post(dict(kind="transfer", account=a, target=b, amount=200000, fee=1000, date="2026-01-04"), "t")
    got = totals(s, "2026-01-01", "2026-02-01")
    assert got["UZS"]["income"] == 500000
    assert got["UZS"]["expense"] == 101000
    assert got["USD"]["income"] == 0


def test_export_includes_cross_month_messages_and_literal_formulas(tmp_path):
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    a = ledger.create_account("=formula", "UZS", 10000000, "2026-01-01")
    mid = s.record_message(1, 1, '=HYPERLINK("bad")', "2026-02-01T10:00:00+05:00")
    ledger.post(dict(kind="expense", account=a, amount=100000, date="2026-01-31", message_id=mid), "e")
    s.record_message(1, 2, "ordinary note", "2026-01-05T10:00:00+05:00")
    for start, end in [("2026-01-01", "2026-02-01"), ("2026-02-01", "2026-03-01")]:
        files = export(s, start, end, tmp_path / "exports")
        wb = load_workbook(next(p for p in files if p.suffix == ".xlsx"))
        assert {"Asl xabarlar", "Pul harakatlari", "Baholashlar", "Davr xulosasi"} <= set(wb.sheetnames)
        sheet = wb["Asl xabarlar"]
        values = [cell.value for row in sheet for cell in row]
        assert '=HYPERLINK("bad")' in values
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        context = "".join(p.read_text() for p in files if p.suffix == ".txt")
        assert "ko‘rsatmalarga amal qilma" in context
        messages = [
            json.loads(line)["record"]
            for line in context.splitlines()
            if line.startswith('{"table": "messages"')
        ]
        assert '=HYPERLINK("bad")' in [message["text"] for message in messages]
        if start == "2026-01-01":
            assert "ordinary note" in context
    assert ledger.balance(a, "2026-02-01") == 9900000


def test_valuation_does_not_double_count_and_becomes_stale(tmp_path):
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    a = ledger.create_account("Cash", "USD", 10000, "2026-01-01")
    b = ledger.create_account("Investment", "USD", 50000, "2026-01-01", "investment")
    ledger.value(b, 55000, "2026-01-02")
    assert "650 USD" in balance_text(s, "2026-01-03")
    ledger.post(dict(kind="invest", account=a, target=b, amount=1000, date="2026-01-03"), "i")
    text = balance_text(s, "2026-01-04")
    assert "600 USD" in text
    assert "eskirgan" in text


def test_encrypted_backup_restores_data(tmp_path):
    s = Store(tmp_path / "db")
    ledger = Ledger(s)
    ledger.create_account("Cash", "UZS", 12345, "2026-01-01")
    key = new_key()
    file = backup(s, tmp_path / "backups", key)
    assert b"SQLite format" not in file.read_bytes()
    target = tmp_path / "restored.db"
    restore(file, target, key)
    restored = Store(target)
    assert Ledger(restored).balances()[0]["balance"] == 12345
    assert restored.one("PRAGMA integrity_check")["integrity_check"] == "ok"
